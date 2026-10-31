"""Locate declared tables, occupied ranges, merged cells and supported drawings."""

import warnings
from pathlib import Path
from io import BytesIO
from stat import S_ISREG
from typing import Any
from zipfile import BadZipFile, ZipFile
from .policy import XlsxLayoutPolicy
from ...models import (
    LayoutAnalysisResult,
    LayoutFinding,
    LayoutRegion,
    LayoutSurface,
    RegionKind,
    XlsxDrawingLocation,
    XlsxRangeLocation,
    XlsxSheetLocation,
)
from ...utils import resolve_source
from ....normalization.artifacts import Artifact


class XlsxLayoutAnalyzer:
    """Return native sheet locations, not printed pages or semantic reading order."""

    def __init__(self, policy: XlsxLayoutPolicy | None = None) -> None:
        self._policy = policy if policy is not None else XlsxLayoutPolicy()

    @property
    def policy(self) -> XlsxLayoutPolicy:
        return self._policy

    @property
    def name(self) -> str:
        return "layout_analysis_xlsx"

    def analyze(self, source: Path | str | bytes | Artifact) -> LayoutAnalysisResult:
        source = resolve_source(source)
        try:
            if source.path is not None:
                path = source.path
                info = path.stat()
                if not S_ISREG(info.st_mode):
                    return self._failure(
                        "source_unreadable", "Source must be a regular file."
                    )
                if info.st_size > self._policy.max_source_bytes:
                    return self._failure(
                        "source_too_large", "Source exceeds the byte limit."
                    )
                with path.open("rb") as stream:
                    data = stream.read(self._policy.max_source_bytes + 1)
            else:
                assert source.data is not None
                data = source.data
        except OSError:
            return self._failure("source_unreadable", "Source could not be read.")
        if len(data) > self._policy.max_source_bytes:
            return self._failure("source_too_large", "Source exceeds the byte limit.")
        try:
            with ZipFile(BytesIO(data)) as archive:
                entries = archive.infolist()
                if (
                    len(entries) > self._policy.max_archive_entries
                    or sum(entry.file_size for entry in entries)
                    > self._policy.max_expanded_bytes
                    or any(entry.flag_bits & 1 for entry in entries)
                ):
                    return self._failure(
                        "archive_limit",
                        "Package exceeds ZIP limits or contains encrypted ZIP entries.",
                    )
                if len({entry.filename for entry in entries}) != len(entries):
                    return self._failure(
                        "decode_failed", "Package contains duplicate part names."
                    )
        except BadZipFile:
            return self._failure(
                "decode_failed", "Source is not a readable workbook package."
            )
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException
        except ImportError as error:
            raise ImportError(
                "Worksheet layout requires: pip install 'raggae[layout-xlsx]'"
            ) from error
        try:
            with warnings.catch_warnings(record=True) as captured:
                warnings.simplefilter("always", UserWarning)
                workbook = load_workbook(
                    BytesIO(data), data_only=False, keep_links=False
                )
        except (
            BadZipFile,
            InvalidFileException,
            SyntaxError,
            ValueError,
            KeyError,
            TypeError,
            OSError,
        ):
            return self._failure(
                "decode_failed", "Source could not be opened as a workbook."
            )
        try:
            result = self._locate(workbook)
        finally:
            workbook.close()
        if captured:
            result = result.model_copy(
                update={
                    "outcome": "partial",
                    "findings": result.findings
                    + (
                        LayoutFinding(
                            code="layout.xlsx.backend_coverage",
                            message="The backend reported unsupported or unreadable workbook features; some content may be missing.",
                        ),
                    ),
                }
            )
        return result

    def _locate(self, workbook: Any) -> LayoutAnalysisResult:
        from openpyxl.worksheet.worksheet import Worksheet

        surfaces, regions, findings = [], [], []
        for number, name in enumerate(workbook.sheetnames, 1):
            if number > self._policy.max_sheets:
                findings.append(
                    LayoutFinding(
                        code="layout.xlsx.sheet_limit",
                        message="Remaining sheets were not analyzed: sheet limit reached.",
                    )
                )
                break
            sheet = workbook[name]
            surface = LayoutSurface(
                id_=f"sheet-{number}",
                kind="sheet",
                location=XlsxSheetLocation(sheet_number=number, sheet_name=name),
            )
            surfaces.append(surface)
            if not isinstance(sheet, Worksheet):
                findings.append(
                    LayoutFinding(
                        code="layout.xlsx.unsupported_sheet",
                        surface_id=surface.id_,
                        message="This non-worksheet sheet was retained but its contents were not mapped.",
                    )
                )
                continue
            self._map_sheet(sheet, surface, regions, findings)
        if not surfaces:
            return self._failure("no_sheets", "Workbook contains no sheets.")
        return LayoutAnalysisResult(
            outcome="partial" if findings else "success",
            surfaces=tuple(surfaces),
            regions=tuple(regions),
            findings=tuple(findings),
        )

    def _map_sheet(
        self,
        sheet: Any,
        surface: LayoutSurface,
        regions: list[LayoutRegion],
        findings: list[LayoutFinding],
    ) -> None:
        from openpyxl.utils.cell import get_column_letter, range_boundaries

        start = len(regions)

        def issue(code: str, message: str) -> None:
            findings.append(
                LayoutFinding(
                    code=f"layout.xlsx.{code}", surface_id=surface.id_, message=message
                )
            )

        def add(
            kind: RegionKind, location: XlsxRangeLocation | XlsxDrawingLocation
        ) -> None:
            if len(regions) - start >= self._policy.max_regions_per_sheet:
                if not any(
                    f.code == "layout.xlsx.region_limit" and f.surface_id == surface.id_
                    for f in findings
                ):
                    issue(
                        "region_limit",
                        "Remaining regions were not returned: region limit reached.",
                    )
                return
            regions.append(
                LayoutRegion(
                    id=f"{surface.id_}-r{len(regions) - start + 1}",
                    surface_id=surface.id_,
                    kind=kind,
                    location=location,
                )
            )

        def reference(box: tuple[int, int, int, int]) -> str:
            left, top, right, bottom = box
            first, last = (
                f"{get_column_letter(left)}{top}",
                f"{get_column_letter(right)}{bottom}",
            )
            return first if first == last else f"{first}:{last}"

        covered = []
        for table in sheet.tables.values():
            covered.append(range_boundaries(table.ref))
            add(
                "table", XlsxRangeLocation(sheet_name=sheet.title, cell_range=table.ref)
            )
        for merged in sorted(
            sheet.merged_cells.ranges, key=lambda item: (item.min_row, item.min_col)
        ):
            if sheet.cell(merged.min_row, merged.min_col).value is None:
                continue
            box = merged.bounds
            if any(
                left <= box[0]
                and top <= box[1]
                and right >= box[2]
                and bottom >= box[3]
                for left, top, right, bottom in covered
            ):
                continue
            covered.append(box)
            add(
                "unknown",
                XlsxRangeLocation(sheet_name=sheet.title, cell_range=str(merged)),
            )

        # Sparse access avoids iterating the enormous rectangle between A1 and
        # a single distant cell. Formatting-only and merged placeholder cells
        # are not occupied content. Formulas are retained, never evaluated.
        occupied = set()
        count = 0
        for (row, column), cell in sorted(sheet._cells.items()):
            if cell.value is None:
                continue
            if count >= self._policy.max_cells_per_sheet:
                issue(
                    "cell_limit",
                    "Remaining occupied cells were not analyzed: cell limit reached.",
                )
                break
            count += 1
            if not any(
                left <= column <= right and top <= row <= bottom
                for left, top, right, bottom in covered
            ):
                occupied.add((row, column))
        for box in self._occupied_blocks(occupied):
            add(
                "unknown",
                XlsxRangeLocation(sheet_name=sheet.title, cell_range=reference(box)),
            )

        for kind, drawings in (("image", sheet._images), ("chart", sheet._charts)):
            for index, drawing in enumerate(drawings):
                marker = getattr(drawing.anchor, "_from", None)
                anchor = (
                    None
                    if marker is None
                    else f"{get_column_letter(marker.col + 1)}{marker.row + 1}"
                )
                location = XlsxDrawingLocation(
                    sheet_name=sheet.title,
                    drawing_kind="image" if kind == "image" else "chart",
                    drawing_index=index,
                    anchor_cell=anchor,
                )
                add("image" if kind == "image" else "unknown", location)
                if kind == "chart":
                    issue(
                        "chart_unanalyzed",
                        "A chart was located; its visual contents were not interpreted.",
                    )
        if (
            sheet.sheet_state != "visible"
            or any(item.hidden for item in sheet.row_dimensions.values())
            or any(item.hidden for item in sheet.column_dimensions.values())
        ):
            issue(
                "hidden_content",
                "Stored content is mapped without applying hidden sheet, row, or column visibility.",
            )
        if sheet._pivots:
            issue(
                "pivot_unanalyzed",
                "Pivot-table semantics and cached data were not interpreted.",
            )

    @staticmethod
    def _occupied_blocks(
        cells: set[tuple[int, int]],
    ) -> list[tuple[int, int, int, int]]:
        """Group four-neighbour cells into (min_col, min_row, max_col, max_row)."""
        remaining = set(cells)
        boxes = []
        for first in sorted(cells):
            if first not in remaining:
                continue
            remaining.remove(first)
            pending = [first]
            top = bottom = first[0]
            left = right = first[1]
            while pending:
                row, column = pending.pop()
                top, bottom = min(top, row), max(bottom, row)
                left, right = min(left, column), max(right, column)
                for neighbour in (
                    (row - 1, column),
                    (row + 1, column),
                    (row, column - 1),
                    (row, column + 1),
                ):
                    if neighbour in remaining:
                        remaining.remove(neighbour)
                        pending.append(neighbour)
            boxes.append((left, top, right, bottom))
        return boxes

    @staticmethod
    def _failure(code: str, message: str) -> LayoutAnalysisResult:
        return LayoutAnalysisResult(
            outcome="failed",
            findings=(LayoutFinding(code=f"layout.xlsx.{code}", message=message),),
        )
