"""Real slide shape locations, group transforms, and ordering compatibility."""

from raggae.documents.layout_analysis import PptxLayoutPolicy

from io import BytesIO
import subprocess
import sys
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from PIL import Image

pptx = pytest.importorskip("pptx")
from pptx.util import Pt
from pptx.enum.shapes import MSO_SHAPE_TYPE

from raggae.documents.layout_analysis import (
    GeometricReadingOrderResolver, LayoutAnalysisResult,
    LayoutAnalyzer, PptxLayoutAnalyzer,
)
from raggae.documents.normalization import Artifact


def encode(presentation):
    with BytesIO() as stream:
        presentation.save(stream)
        return stream.getvalue()


def deck():
    result = pptx.Presentation()
    result.slide_width, result.slide_height = Pt(720), Pt(540)
    return result


def text(shapes, value, x=40, y=60, width=200, height=40):
    shape = shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
    shape.text = value
    return shape


def sample():
    presentation = deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Shipment summary"
    text(slide.shapes, "An ordinary paragraph", y=120)
    slide.shapes.add_table(2, 2, Pt(40), Pt(200), Pt(300), Pt(100))
    image = BytesIO()
    Image.new("RGB", (12, 12), "blue").save(image, format="PNG")
    image.seek(0)
    slide.shapes.add_picture(image, Pt(400), Pt(200), Pt(100), Pt(80))
    presentation.slides.add_slide(presentation.slide_layouts[6])
    return presentation


@pytest.mark.parametrize("source_kind", ["bytes", "path", "string", "artifact_bytes", "artifact_path"])
@pytest.mark.parametrize("positional_source", [False, True])
def test_public_inputs_share_map_without_changing_source(tmp_path, source_kind, positional_source):
    data = encode(sample())
    path = tmp_path / "slides.bin"
    path.write_bytes(data)
    source = {"bytes": data, "path": path, "string": str(path), "artifact_bytes": Artifact(data=data),
              "artifact_path": Artifact(path=path)}[source_kind]
    worker = PptxLayoutAnalyzer()
    result = worker.analyze(source) if positional_source else worker.analyze(source=source)
    assert result.outcome == "success"
    assert [region.kind for region in result.regions] == ["heading", "text", "table", "image"]
    assert len(result.surfaces) == 2 and all(surface.kind == "slide" for surface in result.surfaces)
    assert (result.surfaces[0].width, result.surfaces[0].height, result.surfaces[0].unit) == (720, 540, "pt")
    assert result.regions[-1].bbox == pytest.approx((400, 200, 500, 280))
    assert all(region.native_order is None for region in result.regions)
    assert "Shipment summary" not in result.model_dump_json()
    assert LayoutAnalysisResult.model_validate_json(result.model_dump_json()) == result
    assert path.read_bytes() == data
    assert isinstance(worker, LayoutAnalyzer)


def test_shape_paths_resolve_to_original_shapes():
    presentation = sample()
    result = PptxLayoutAnalyzer().analyze(source=encode(presentation))
    for region in result.regions:
        slide = next(slide for slide in presentation.slides if str(slide.part.partname).lstrip("/") == region.location.part)
        shapes = slide.shapes
        for index in region.location.shape_path:
            shape = shapes[index]
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                shapes = shape.shapes
        assert shape.shape_id == region.location.shape_id


@pytest.mark.parametrize("flip_h", [False, True])
@pytest.mark.parametrize("flip_v", [False, True])
def test_group_translation_scaling_and_mirroring(flip_h, flip_v):
    presentation = deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    text(group.shapes, "First", 40, 60, 100, 30)
    text(group.shapes, "Second", 180, 120, 100, 30)
    group.left, group.top, group.width, group.height = map(Pt, (100, 200, 480, 180))
    group.element.xfrm.flipH, group.element.xfrm.flipV = flip_h, flip_v
    result = PptxLayoutAnalyzer().analyze(source=encode(presentation))
    assert result.outcome == "success"
    first, second = result.regions
    assert first.bbox == pytest.approx((380 if flip_h else 100, 320 if flip_v else 200,
                                        580 if flip_h else 300, 380 if flip_v else 260))
    assert second.bbox == pytest.approx((100 if flip_h else 380, 200 if flip_v else 320,
                                         300 if flip_h else 580, 260 if flip_v else 380))
    assert first.location.shape_path == (0, 0)


def test_group_depth_policy_can_stop_before_visiting_children():
    presentation = deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    text(group.shapes, "Child", 40, 60, 100, 30)
    data = encode(presentation)
    assert PptxLayoutAnalyzer().analyze(source=data).outcome == "success"
    limited = PptxLayoutAnalyzer(policy=PptxLayoutPolicy(max_group_depth=1)).analyze(source=data)
    assert limited.outcome == "partial"
    assert len(limited.regions) == 1 and limited.regions[0].kind == "unknown"
    assert limited.regions[0].location.shape_path == (0,)
    assert limited.findings[0].code == "layout.pptx.group_transform_unresolved"


def test_nested_groups_compose_parent_and_child_transforms():
    presentation = deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    text(inner.shapes, "Nested", 10, 20, 40, 30)
    inner.left, inner.top, inner.width, inner.height = map(Pt, (100, 50, 80, 60))
    xfrm = outer.element.xfrm
    xfrm.chOff.x, xfrm.chOff.y = 0, 0
    xfrm.chExt.cx, xfrm.chExt.cy = Pt(200), Pt(150)
    outer.left, outer.top, outer.width, outer.height = map(Pt, (50, 60, 400, 300))
    result = PptxLayoutAnalyzer().analyze(source=encode(presentation))
    assert result.outcome == "success"
    assert result.regions[0].bbox == pytest.approx((250, 160, 410, 280))
    assert result.regions[0].location.shape_path == (0, 0, 0)


@pytest.mark.parametrize("grouped", [False, True])
def test_rotations_preserve_locator_without_misleading_box(grouped):
    presentation = deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape() if grouped else None
    shape = text(group.shapes if group is not None else slide.shapes, "Rotated")
    (group if group is not None else shape).rotation = 30
    result = PptxLayoutAnalyzer().analyze(source=encode(presentation))
    assert result.outcome == "partial"
    assert result.regions[0].bbox is None and result.regions[0].location.shape_id > 0
    assert result.findings[0].code in ("layout.pptx.group_transform_unresolved", "layout.pptx.shape_geometry_unresolved")
    assert GeometricReadingOrderResolver().resolve_order(result).unresolved_region_ids


def test_clipping_and_fully_offslide_shapes():
    presentation = deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text(slide.shapes, "Partial", -20, 60, 100, 30)
    text(slide.shapes, "Outside", 800, 60, 100, 30)
    result = PptxLayoutAnalyzer().analyze(source=encode(presentation))
    assert result.outcome == "partial"
    assert result.regions[0].bbox == pytest.approx((0, 60, 80, 90))
    assert result.regions[1].bbox is None


def test_storage_order_is_not_reading_order_and_geometry_can_follow():
    presentation = deck()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text(slide.shapes, "Bottom created first", y=300)
    text(slide.shapes, "Top created last", y=60)
    layout = PptxLayoutAnalyzer().analyze(source=encode(presentation))
    order = GeometricReadingOrderResolver().resolve_order(layout)
    assert order.outcome == "success"
    assert order.edges[0].before_region_id == layout.regions[1].id


def test_slide_sequence_is_presentation_order_not_filename_order():
    presentation = sample()
    slides = presentation.slides._sldIdLst
    first = slides[0]
    slides.remove(first)
    slides.append(first)
    data = encode(presentation)
    result = PptxLayoutAnalyzer().analyze(source=data)
    assert result.surfaces[0].location.part.endswith("slide2.xml")
    assert result.surfaces[0].location.slide_number == 1
    assert not any(region.surface_id == result.surfaces[0].id_ for region in result.regions)
    assert result.surfaces[1].location.part.endswith("slide1.xml")
    with ZipFile(BytesIO(data)) as archive:
        for region in result.regions:
            element = ElementTree.fromstring(archive.read(region.location.part))
            shapes = element.find("{http://schemas.openxmlformats.org/presentationml/2006/main}cSld").find(
                "{http://schemas.openxmlformats.org/presentationml/2006/main}spTree")
            # The first two children describe the group itself, not slide shapes.
            shape = list(shapes)[region.location.shape_path[0] + 2]
            identifier = shape.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}cNvPr")
            assert int(identifier.get("id")) == region.location.shape_id


def test_limits_and_speaker_notes():
    presentation = sample()
    presentation.slides[0].notes_slide.notes_text_frame.text = "Example speaker note"
    data = encode(presentation)
    assert PptxLayoutAnalyzer(policy=PptxLayoutPolicy(max_source_bytes=len(data)-1)).analyze(source=data).outcome == "failed"
    assert PptxLayoutAnalyzer(policy=PptxLayoutPolicy(max_expanded_bytes=10)).analyze(source=data).findings[0].code == "layout.pptx.archive_limit"
    assert PptxLayoutAnalyzer(policy=PptxLayoutPolicy(max_archive_entries=1)).analyze(source=data).findings[0].code == "layout.pptx.archive_limit"
    assert len(PptxLayoutAnalyzer(policy=PptxLayoutPolicy(max_slides=1)).analyze(source=data).surfaces) == 1
    result = PptxLayoutAnalyzer(policy=PptxLayoutPolicy(max_shapes_per_slide=1)).analyze(source=data)
    assert result.outcome == "partial"
    assert {finding.code for finding in result.findings} >= {"layout.pptx.shape_limit", "layout.pptx.notes_unanalyzed"}


@pytest.mark.parametrize("data", [b"", b"not a presentation", b"PK malformed"])
def test_bad_packages_fail_safely(data):
    result = PptxLayoutAnalyzer().analyze(source=data)
    assert result.outcome == "failed" and not result.surfaces


def test_empty_presentation_and_blank_slide():
    presentation = deck()
    assert PptxLayoutAnalyzer().analyze(source=encode(presentation)).outcome == "failed"
    presentation.slides.add_slide(presentation.slide_layouts[6])
    result = PptxLayoutAnalyzer().analyze(source=encode(presentation))
    assert result.outcome == "success" and not result.regions


def test_optional_backend_not_needed_for_public_import():
    code = """
import builtins
original = builtins.__import__
def guarded(name, globals=None, locals=None, fromlist=(), level=0):
    if level == 0 and name.split('.')[0] == 'pptx':
        raise ImportError('external backend unavailable')
    return original(name, globals, locals, fromlist, level)
builtins.__import__ = guarded
from raggae.documents.layout_analysis import PptxLayoutAnalyzer
assert PptxLayoutAnalyzer().name == 'layout_analysis_pptx'
"""
    result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
