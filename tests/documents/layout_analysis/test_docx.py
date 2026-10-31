"""Real DOCX structure, resolvable element locations, and shared-map compatibility."""

from raggae.documents.layout_analysis import DocxLayoutPolicy

from io import BytesIO
import subprocess
import sys
from xml.etree import ElementTree
from zipfile import ZipFile

import pytest
from PIL import Image
from pydantic import ValidationError

docx = pytest.importorskip("docx")
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from raggae.documents.layout_analysis import (
    DocxBodyLocation, DocxElementLocation, DocxLayoutAnalyzer, GeometricReadingOrderResolver,
    LayoutAnalysisResult, LayoutAnalyzer, LayoutRegion, LayoutSurface,
)
from raggae.documents.normalization import Artifact


def picture():
    stream = BytesIO()
    Image.new("RGB", (24, 12), "blue").save(stream, format="PNG")
    stream.seek(0)
    return stream


def encode(document):
    with BytesIO() as stream:
        document.save(stream)
        return stream.getvalue()


def fixture_document():
    document = docx.Document()
    document.add_heading("Operations report", level=0)
    document.add_paragraph("Example introductory paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Count"
    table.cell(1, 0).text = "Cartons"
    table.cell(1, 1).text = "12"
    document.add_picture(picture())
    document.add_paragraph("Example closing paragraph.")
    return document


@pytest.mark.parametrize("source_kind", ["bytes", "path", "string", "artifact_bytes", "artifact_path"])
@pytest.mark.parametrize("positional_source", [False, True])
def test_same_map_for_all_public_inputs(source_kind, positional_source, tmp_path):
    data = encode(fixture_document())
    path = tmp_path / "report.bin"
    path.write_bytes(data)
    source = {"bytes": data, "path": path, "string": str(path),
              "artifact_bytes": Artifact(data=data), "artifact_path": Artifact(path=path)}[source_kind]
    worker = DocxLayoutAnalyzer()
    result = worker.analyze(source) if positional_source else worker.analyze(source=source)
    assert result.outcome == "success"
    assert [region.kind for region in result.regions] == ["heading", "text", "table", "image", "text"]
    assert [region.native_order for region in result.regions] == list(range(5))
    assert all(region.bbox is None and region.location is not None for region in result.regions)
    surface, = result.surfaces
    assert surface.kind == "flow" and surface.width is surface.height is surface.unit is None
    assert "Example introductory" not in result.model_dump_json()
    assert LayoutAnalysisResult.model_validate_json(result.model_dump_json()) == result
    assert path.read_bytes() == data
    assert isinstance(worker, LayoutAnalyzer)


def test_locators_resolve_to_actual_source_elements():
    data = encode(fixture_document())
    result = DocxLayoutAnalyzer().analyze(source=data)
    with ZipFile(BytesIO(data)) as archive:
        for region in result.regions:
            locator = region.location
            body = ElementTree.fromstring(archive.read(locator.part)).find(qn("w:body"))
            element = body
            for index in locator.child_indexes:
                element = list(element)[index]
            expected = "w:tbl" if region.kind == "table" else "w:drawing" if region.kind == "image" else "w:p"
            assert element.tag == qn(expected)


def test_nested_table_and_picture_do_not_duplicate_cell_text():
    document = docx.Document()
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.text = "This stays inside the table"
    cell.add_table(rows=1, cols=1).cell(0, 0).text = "Nested cell"
    cell.add_paragraph().add_run().add_picture(picture())
    result = DocxLayoutAnalyzer().analyze(source=encode(document))
    assert [region.kind for region in result.regions] == ["table", "image"]
    assert len(result.regions[1].location.child_indexes) > 1


def test_styled_heading_is_distinct_from_bold_body_text():
    document = docx.Document()
    style = document.styles.add_style("Local heading", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles["Heading 2"]
    document.add_paragraph("Derived heading", style)
    document.add_paragraph().add_run("Just bold text").bold = True
    document.add_paragraph("  ")
    document.add_page_break()
    result = DocxLayoutAnalyzer().analyze(source=encode(document))
    assert [region.kind for region in result.regions] == ["heading", "text"]


def test_headers_and_content_controls_report_missing_coverage():
    document = docx.Document()
    document.add_paragraph("Body")
    document.sections[0].header.paragraphs[0].text = "Header"
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    content.append(paragraph)
    control.append(content)
    document.element.body.insert(1, control)
    result = DocxLayoutAnalyzer().analyze(source=encode(document))
    assert result.outcome == "partial"
    assert [region.kind for region in result.regions] == ["text", "unknown"]
    assert {finding.code for finding in result.findings} >= {"layout.docx.other_stories", "layout.docx.unsupported_block"}


def test_floating_image_anchor_is_located_without_claiming_page_geometry():
    document = docx.Document()
    document.add_picture(picture())
    inline = next(document.element.body.iter(qn("wp:inline")))
    inline.tag = qn("wp:anchor")
    result = DocxLayoutAnalyzer().analyze(source=encode(document))
    assert result.outcome == "partial"
    assert result.regions[0].kind == "image" and result.regions[0].bbox is None
    assert result.findings[0].code == "layout.docx.drawing_partial"


def test_geometric_worker_declines_unrendered_word_flow():
    layout = DocxLayoutAnalyzer().analyze(source=encode(fixture_document()))
    result = GeometricReadingOrderResolver().resolve_order(layout)
    assert result.outcome == "not_applicable" and not result.edges
    assert result.findings[-1].code == "layout.order.no_geometry"


def test_region_limit_and_archive_limit_are_explicit():
    data = encode(fixture_document())
    partial = DocxLayoutAnalyzer(policy=DocxLayoutPolicy(max_regions=2)).analyze(source=data)
    assert partial.outcome == "partial" and len(partial.regions) == 2
    assert partial.findings[0].code == "layout.docx.region_limit"
    assert DocxLayoutAnalyzer(policy=DocxLayoutPolicy(max_regions=5)).analyze(source=data).outcome == "success"
    assert DocxLayoutAnalyzer(policy=DocxLayoutPolicy(max_source_bytes=len(data) - 1)).analyze(source=data).outcome == "failed"
    assert DocxLayoutAnalyzer(policy=DocxLayoutPolicy(max_expanded_bytes=10)).analyze(source=data).findings[0].code == "layout.docx.archive_limit"
    assert DocxLayoutAnalyzer(policy=DocxLayoutPolicy(max_archive_entries=1)).analyze(source=data).findings[0].code == "layout.docx.archive_limit"


@pytest.mark.parametrize("data", [b"", b"not zip", b"PK broken archive"])
def test_invalid_packages_return_safe_findings(data):
    result = DocxLayoutAnalyzer().analyze(source=data)
    assert result.outcome == "failed" and not result.regions and not result.surfaces
    assert result.findings[0].code == "layout.docx.decode_failed"


def test_blank_document_is_a_valid_empty_flow():
    result = DocxLayoutAnalyzer().analyze(source=encode(docx.Document()))
    assert result.outcome == "success" and len(result.surfaces) == 1 and result.regions == ()


def test_public_import_without_external_docx_backend():
    code = """
import builtins
original = builtins.__import__
def guarded(name, globals=None, locals=None, fromlist=(), level=0):
    if level == 0 and name.split('.')[0] == 'docx':
        raise ImportError('optional backend unavailable')
    return original(name, globals, locals, fromlist, level)
builtins.__import__ = guarded
from raggae.documents.layout_analysis import DocxLayoutAnalyzer
assert DocxLayoutAnalyzer().name == 'layout_analysis_docx'
"""
    result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


def test_flow_cannot_invent_geometry_and_requires_a_location():
    with pytest.raises(ValidationError):
        LayoutSurface(id_="body", kind="flow", location=DocxBodyLocation(part="word/document.xml"), width=600)
    with pytest.raises(ValidationError):
        LayoutRegion(id="r1", surface_id="body", kind="text")
    location = DocxElementLocation(part="word/document.xml", child_indexes=(0,))
    assert LayoutRegion(id="r1", surface_id="body", kind="text", location=location).bbox is None


@pytest.mark.parametrize("name", ["max_source_bytes", "max_expanded_bytes", "max_regions"])
@pytest.mark.parametrize("value", [True, 0, -1, "large"])
def test_bad_limits_rejected(name, value):
    with pytest.raises((TypeError, ValueError)):
        DocxLayoutAnalyzer(policy=DocxLayoutPolicy(**{name: value}))
