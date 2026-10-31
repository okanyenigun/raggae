"""Real worksheet locations, sparse-block boundaries, and honest coverage."""

from raggae.documents.layout_analysis import XlsxLayoutPolicy

from io import BytesIO
import subprocess
import sys
import warnings
from zipfile import ZipFile

import pytest
from PIL import Image
from pydantic import ValidationError

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.drawing.spreadsheet_drawing import AbsoluteAnchor
from openpyxl.styles import PatternFill
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.table import Table

from raggae.documents.layout_analysis import (
    GeometricReadingOrderResolver, LayoutAnalysisResult,
    LayoutAnalyzer, LayoutRegion, LayoutSurface, XlsxDrawingLocation, XlsxLayoutAnalyzer,
    XlsxRangeLocation, XlsxSheetLocation,
)
from raggae.documents.normalization import Artifact


def encode(workbook):
    with BytesIO() as output:
        workbook.save(output)
        return output.getvalue()


def sample():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Dispatch & costs"
    sheet.append(["Item", "Count"])
    sheet.append(["Box", 2])
    sheet.append(["Crate", 4])
    sheet.add_table(Table(displayName="Dispatch", ref="A1:B3"))
    sheet["F2"], sheet["G2"], sheet["F3"], sheet["G3"] = "Cost", "Total", 5, "=SUM(B2:B3)"
    sheet["A6"] = "A merged label"
    sheet.merge_cells("A6:D6")
    sheet["XFD1048576"].fill = PatternFill("solid", fgColor="FF0000")
    workbook.create_sheet("Empty")
    return workbook


@pytest.mark.parametrize("source_kind", ["bytes", "path", "string", "artifact_bytes", "artifact_path"])
@pytest.mark.parametrize("positional_source", [False, True])
def test_public_input_styles_and_native_ranges(tmp_path, source_kind, positional_source):
    data = encode(sample())
    path = tmp_path / "workbook.bin"
    path.write_bytes(data)
    source = {"bytes": data, "path": path, "string": str(path), "artifact_bytes": Artifact(data=data),
              "artifact_path": Artifact(path=path)}[source_kind]
    worker = XlsxLayoutAnalyzer()
    result = worker.analyze(source) if positional_source else worker.analyze(source=source)
    assert result.outcome == "success"
    assert isinstance(worker, LayoutAnalyzer) and worker.name == "layout_analysis_xlsx"
    assert len(result.surfaces) == 2
    assert result.surfaces[0].location.sheet_name == "Dispatch & costs"
    assert all(surface.kind == "sheet" and surface.width is None and surface.unit is None for surface in result.surfaces)
    assert [(region.kind, region.location.cell_range) for region in result.regions] == [
        ("table", "A1:B3"), ("unknown", "A6:D6"), ("unknown", "F2:G3"),
    ]
    assert all(region.bbox is None and region.native_order is None for region in result.regions)
    assert "A merged label" not in result.model_dump_json()
    assert LayoutAnalysisResult.model_validate_json(result.model_dump_json()) == result
    assert path.read_bytes() == data
    assert GeometricReadingOrderResolver().resolve_order(result).outcome == "not_applicable"


def test_ranges_resolve_to_original_cells():
    data = encode(sample())
    result = XlsxLayoutAnalyzer().analyze(source=data)
    workbook = openpyxl.load_workbook(BytesIO(data), data_only=False)
    try:
        cells = {region.location.cell_range: workbook[region.location.sheet_name][region.location.cell_range] for region in result.regions}
        assert cells["A1:B3"][1][0].value == "Box"
        assert cells["F2:G3"][1][1].value == "=SUM(B2:B3)"
        assert cells["A6:D6"][0][0].value == "A merged label"
    finally:
        workbook.close()


def test_separate_blocks_adjacency_holes_and_diagonals():
    cells = {(1, 1), (1, 2), (2, 1), (1, 5), (2, 5), (4, 4), (5, 5)}
    assert XlsxLayoutAnalyzer._occupied_blocks(cells) == [(1, 1, 2, 2), (5, 1, 5, 2), (4, 4, 4, 4), (5, 5, 5, 5)]
    assert cells == {(1, 1), (1, 2), (2, 1), (1, 5), (2, 5), (4, 4), (5, 5)}
    assert XlsxLayoutAnalyzer._occupied_blocks(set()) == []


def test_sparse_distant_cells_and_zero_are_not_skipped_or_expanded():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet["A1"] = 0
    sheet["XFD1048576"] = False
    result = XlsxLayoutAnalyzer().analyze(source=encode(workbook))
    assert result.outcome == "success"
    assert [region.location.cell_range for region in result.regions] == ["A1", "XFD1048576"]


def test_images_charts_and_absolute_anchor_have_resolvable_native_locations():
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Value"])
    sheet.append([5])
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=1, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "F2")
    picture = BytesIO()
    Image.new("RGB", (10, 10), "blue").save(picture, format="PNG")
    picture.seek(0)
    sheet.add_image(WorksheetImage(picture), "C4")
    second_picture = BytesIO(picture.getvalue())
    absolute = WorksheetImage(second_picture)
    absolute.anchor = AbsoluteAnchor()
    sheet.add_image(absolute)
    data = encode(workbook)
    result = XlsxLayoutAnalyzer().analyze(source=data)
    assert result.outcome == "partial"
    drawings = [region for region in result.regions if isinstance(region.location, XlsxDrawingLocation)]
    original = openpyxl.load_workbook(BytesIO(data))
    try:
        # The backend groups drawings by anchor type on reload. Locations index
        # those reloaded collections, not the fixture builder's insertion order.
        expected = []
        for kind, collection in (("image", original.active._images), ("chart", original.active._charts)):
            for index, drawing in enumerate(collection):
                marker = getattr(drawing.anchor, "_from", None)
                anchor = None if marker is None else f"{get_column_letter(marker.col + 1)}{marker.row + 1}"
                expected.append(("image" if kind == "image" else "unknown", kind, index, anchor))
        assert [(region.kind, region.location.drawing_kind, region.location.drawing_index, region.location.anchor_cell)
                for region in drawings] == expected
        assert {region.location.anchor_cell for region in drawings if region.kind == "image"} == {None, "C4"}
        for region in drawings:
            location = region.location
            collection = original[location.sheet_name]._images if location.drawing_kind == "image" else original[location.sheet_name]._charts
            drawing = collection[location.drawing_index]
            if location.anchor_cell is not None:
                assert drawing.anchor._from.col + 1 == (3 if location.drawing_kind == "image" else 6)
    finally:
        original.close()
    assert any(finding.code == "layout.xlsx.chart_unanalyzed" for finding in result.findings)


def test_hidden_sheets_rows_and_columns_are_retained_with_findings():
    workbook = sample()
    sheet = workbook.active
    sheet.row_dimensions[2].hidden = True
    sheet.column_dimensions["F"].hidden = True
    workbook["Empty"].sheet_state = "hidden"
    result = XlsxLayoutAnalyzer().analyze(source=encode(workbook))
    assert result.outcome == "partial" and len(result.surfaces) == 2
    assert {finding.surface_id for finding in result.findings if finding.code == "layout.xlsx.hidden_content"} == {"sheet-1", "sheet-2"}


def test_sheet_order_uses_workbook_order_and_exact_names():
    workbook = sample()
    workbook.move_sheet("Empty", offset=-1)
    result = XlsxLayoutAnalyzer().analyze(source=encode(workbook))
    assert [surface.location.sheet_name for surface in result.surfaces] == ["Empty", "Dispatch & costs"]
    assert result.surfaces[0].location.sheet_number == 1
    assert all(region.surface_id == "sheet-2" for region in result.regions)


def test_limits_produce_bounded_maps_and_safe_failures():
    data = encode(sample())
    assert XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(max_source_bytes=len(data)-1)).analyze(source=data).findings[0].code == "layout.xlsx.source_too_large"
    assert XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(max_source_bytes=len(data))).analyze(source=data).outcome == "success"
    assert XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(max_expanded_bytes=10)).analyze(source=data).findings[0].code == "layout.xlsx.archive_limit"
    assert XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(max_archive_entries=1)).analyze(source=data).findings[0].code == "layout.xlsx.archive_limit"
    result = XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(max_sheets=1)).analyze(source=data)
    assert len(result.surfaces) == 1 and result.outcome == "partial"
    result = XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(max_regions_per_sheet=1)).analyze(source=data)
    assert len(result.regions) == 1 and result.outcome == "partial"
    assert sum(finding.code == "layout.xlsx.region_limit" for finding in result.findings) == 1
    result = XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(max_cells_per_sheet=1)).analyze(source=data)
    assert result.outcome == "partial" and any(finding.code == "layout.xlsx.cell_limit" for finding in result.findings)


@pytest.mark.parametrize("name", ["max_source_bytes", "max_expanded_bytes", "max_sheets", "max_cells_per_sheet", "max_regions_per_sheet"])
@pytest.mark.parametrize("value, error", [(True, ValidationError), (1.5, ValidationError), (0, ValidationError), (-1, ValidationError)])
def test_constructor_limits(name, value, error):
    with pytest.raises(error):
        XlsxLayoutAnalyzer(policy=XlsxLayoutPolicy(**{name: value}))


@pytest.mark.parametrize("data", [b"", b"not a workbook", b"PK malformed"])
def test_bad_package(data):
    result = XlsxLayoutAnalyzer().analyze(source=data)
    assert result.outcome == "failed" and not result.surfaces


def test_wrong_office_package_and_missing_source(tmp_path):
    with BytesIO() as stream:
        with ZipFile(stream, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
        assert XlsxLayoutAnalyzer().analyze(source=stream.getvalue()).outcome == "failed"
    assert XlsxLayoutAnalyzer().analyze(source=tmp_path / "missing.xlsx").outcome == "failed"
    assert XlsxLayoutAnalyzer().analyze(source=tmp_path).outcome == "failed"


def test_blank_workbook_is_valid():
    result = XlsxLayoutAnalyzer().analyze(source=encode(openpyxl.Workbook()))
    assert result.outcome == "success" and len(result.surfaces) == 1 and not result.regions


def test_backend_warning_marks_coverage_without_leaking_diagnostics(monkeypatch):
    workbook = sample()
    data = encode(workbook)

    def load(*args, **kwargs):
        warnings.warn("private backend details", UserWarning)
        return workbook

    monkeypatch.setattr(openpyxl, "load_workbook", load)
    result = XlsxLayoutAnalyzer().analyze(source=data)
    assert result.outcome == "partial"
    assert result.findings[-1].code == "layout.xlsx.backend_coverage"
    assert "private backend details" not in result.model_dump_json()


def test_optional_backend_not_needed_for_public_import():
    code = """
import builtins
original = builtins.__import__
def guarded(name, globals=None, locals=None, fromlist=(), level=0):
    if level == 0 and name.split('.')[0] == 'openpyxl':
        raise ImportError('external backend unavailable')
    return original(name, globals, locals, fromlist, level)
builtins.__import__ = guarded
from raggae.documents.layout_analysis import XlsxLayoutAnalyzer
assert XlsxLayoutAnalyzer().name == 'layout_analysis_xlsx'
"""
    result = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr


@pytest.mark.parametrize("reference", ["A0", "a1", "XFE1", "A1048577", "B2:A1", "A1:B2:C3", "A1:"])
def test_invalid_range_models(reference):
    with pytest.raises(ValidationError):
        XlsxRangeLocation(sheet_name="Sheet", cell_range=reference)


def test_sheet_models_reject_geometry_and_cross_sheet_regions():
    location = XlsxSheetLocation(sheet_number=1, sheet_name="Sheet")
    with pytest.raises(ValidationError):
        LayoutSurface(id_="s1", kind="sheet", location=location, width=10)
    surface = LayoutSurface(id_="s1", kind="sheet", location=location)
    region = LayoutRegion(id="r1", surface_id="s1", kind="unknown", location=XlsxRangeLocation(sheet_name="Other", cell_range="A1"))
    with pytest.raises(ValidationError):
        LayoutAnalysisResult(outcome="success", surfaces=(surface,), regions=(region,))
