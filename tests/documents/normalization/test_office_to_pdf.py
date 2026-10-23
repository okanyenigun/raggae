"""Office-to-PDF routes, private native profiles, and failure cleanup."""

import shutil
import subprocess
from io import BytesIO
from pathlib import Path

import pytest
from pydantic import ValidationError
from pypdf import PdfReader

from raggae.documents.normalization import (
    Artifact,
    LibreOfficeToPdfConverter,
    NormalizationRequest,
    NormalizationResult,
    NormalizationWorker,
    OfficeToPdfFinding,
    OfficeToPdfIssue,
    OfficeToPdfReport,
)
from raggae.documents.normalization.format_conversion import libreoffice

SOURCE = b"input bytes are opaque to the native-tool boundary"
PDF = b"%PDF-1.7\nmock native output\n%%EOF\n"


def native_paths(args):
    source = Path(args[-1])
    directory = Path(args[args.index("--outdir") + 1])
    return source, directory / "source.pdf"


@pytest.fixture
def fake_engine(monkeypatch):
    calls = []
    monkeypatch.setattr(libreoffice.shutil, "which", lambda value: "/fake tools/soffice")

    def run(args, **kwargs):
        source, output = native_paths(args)
        profile = source.parent / "profile"
        assert profile.is_dir()
        assert list(profile.iterdir()) == []
        assert f"-env:UserInstallation={profile.resolve().as_uri()}" in args
        (profile / "native-settings").write_bytes(b"private profile state")
        calls.append((args, kwargs, source.read_bytes()))
        output.write_bytes(PDF)
        return subprocess.CompletedProcess(args, 0)

    monkeypatch.setattr(libreoffice.subprocess, "run", run)
    return calls


@pytest.mark.parametrize("source_format", ["docx", "xlsx", "pptx"])
@pytest.mark.parametrize("source_kind", ["bytes", "path", "artifact"])
@pytest.mark.parametrize("disk", [False, True])
def test_routes_share_source_and_storage_contract(source_format, source_kind, disk, fake_engine, tmp_path):
    source = SOURCE
    if source_kind == "path":
        source = tmp_path / "original; wrong-extension.bin"
        source.write_bytes(SOURCE)
    elif source_kind == "artifact":
        source = Artifact(data=SOURCE)
    destination = tmp_path / "result.pdf" if disk else None
    request = NormalizationRequest(source=source, output_path=destination)
    worker = LibreOfficeToPdfConverter(source_format, timeout_seconds=17)
    result = worker.normalize(request)
    assert type(result) is NormalizationResult
    assert result.outcome == "changed"
    assert result.report.source_format == source_format
    assert result.report.output_size_bytes == len(PDF)
    assert result.report.issues == ()
    assert result.artifact.read_bytes() == PDF
    assert result.artifact.path == destination
    assert request.source.read_bytes() == SOURCE
    assert OfficeToPdfReport.model_validate_json(result.report.model_dump_json()) == result.report
    args, options, staged_bytes = fake_engine[0]
    staged, candidate = native_paths(args)
    assert staged_bytes == SOURCE
    assert staged.name == f"source.{source_format}"
    assert staged.parent == options["cwd"]
    assert candidate.parent == staged.parent / "output"
    assert args[0] == "/fake tools/soffice"
    assert args[args.index("--convert-to") + 1] == "pdf"
    assert all(flag in args for flag in ("--headless", "--nologo", "--nodefault", "--norestore"))
    assert options["shell"] is False
    assert options["check"] is False
    assert options["timeout"] == 17
    assert options["stdin"] == options["stdout"] == options["stderr"] == subprocess.DEVNULL
    assert not staged.parent.exists()
    assert not list(tmp_path.glob(".raggae-*.part"))


@pytest.mark.parametrize("available,expected_probes", [
    ("soffice", ["soffice"]),
    ("libreoffice", ["soffice", "libreoffice"]),
])
def test_default_executable_discovery_order(available, expected_probes, fake_engine, monkeypatch):
    probes = []

    def which(candidate):
        probes.append(candidate)
        return f"/tools/{candidate}" if candidate == available else None

    monkeypatch.setattr(libreoffice.shutil, "which", which)
    result = LibreOfficeToPdfConverter().normalize(NormalizationRequest(source=SOURCE, output_path=None))
    assert result.outcome == "changed"
    assert probes == expected_probes
    assert fake_engine[0][0][0] == f"/tools/{available}"
    assert fake_engine[0][1]["timeout"] == 120


@pytest.mark.parametrize("executable", [None, "custom-office"])
def test_missing_dependency_does_not_run_or_publish(executable, monkeypatch, tmp_path):
    probes = []

    def which(candidate):
        probes.append(candidate)
        return None

    def forbidden(*args, **kwargs):
        pytest.fail("a missing executable must not be run")

    monkeypatch.setattr(libreoffice.shutil, "which", which)
    monkeypatch.setattr(libreoffice.subprocess, "run", forbidden)
    result = LibreOfficeToPdfConverter(executable=executable).normalize(
        NormalizationRequest(source=SOURCE, output_path=tmp_path / "result.pdf")
    )
    assert probes == (["custom-office"] if executable else ["soffice", "libreoffice"])
    assert result.artifact is None
    assert result.outcome == "failed"
    assert result.report.output_size_bytes is None
    assert result.report.issues[0].code == OfficeToPdfFinding.DEPENDENCY_UNAVAILABLE
    assert not list(tmp_path.iterdir())


def test_custom_relative_executable_made_absolute(fake_engine, monkeypatch):
    probes = []

    def which(candidate):
        probes.append(candidate)
        return "./custom tools/soffice"

    monkeypatch.setattr(libreoffice.shutil, "which", which)
    result = LibreOfficeToPdfConverter(executable=Path("custom tools/soffice")).normalize(
        NormalizationRequest(source=SOURCE, output_path=None)
    )
    assert result.outcome == "changed"
    assert probes == ["custom tools/soffice"]
    assert fake_engine[0][0][0] == str(Path("custom tools/soffice").absolute())


@pytest.mark.parametrize("failure", ["timeout", "start", "nonzero", "missing", "wrong_name", "directory", "empty", "non_pdf", "symlink"])
def test_native_failures_safe_and_both_scratch_areas_cleaned(failure, fake_engine, monkeypatch, tmp_path):
    scratch = []
    external = tmp_path / "private-source"
    external.write_bytes(PDF)

    def run(args, **kwargs):
        source, output = native_paths(args)
        scratch.append(source.parent)
        (source.parent / "profile" / "private-settings").write_bytes(b"private")
        if failure == "timeout":
            output.write_bytes(b"partial output")
            raise subprocess.TimeoutExpired("private command", 17, stderr=b"private source text")
        if failure == "start":
            raise OSError("private native diagnostic")
        if failure == "nonzero":
            output.write_bytes(PDF)
            return subprocess.CompletedProcess(args, 42)
        if failure == "wrong_name":
            (output.parent / "other.pdf").write_bytes(PDF)
        elif failure == "directory":
            output.mkdir()
        elif failure == "empty":
            output.write_bytes(b"")
        elif failure == "non_pdf":
            output.write_bytes(b"<html>not a PDF</html>")
        elif failure == "symlink":
            output.symlink_to(external)
        return subprocess.CompletedProcess(args, 0)

    monkeypatch.setattr(libreoffice.subprocess, "run", run)
    result = LibreOfficeToPdfConverter("pptx").normalize(
        NormalizationRequest(source=SOURCE, output_path=tmp_path / "result.pdf")
    )
    expected = {
        "timeout": OfficeToPdfFinding.TIMEOUT,
        "start": OfficeToPdfFinding.DEPENDENCY_UNAVAILABLE,
        "nonzero": OfficeToPdfFinding.CONVERSION_FAILED,
    }.get(failure, OfficeToPdfFinding.INVALID_OUTPUT)
    assert result.artifact is None
    assert result.report.issues[0].code == expected
    assert result.report.source_format == "pptx"
    assert result.report.output_size_bytes is None
    assert "private" not in result.report.model_dump_json()
    assert all(not path.exists() for path in scratch)
    assert not (tmp_path / "result.pdf").exists()
    assert not list(tmp_path.glob(".raggae-*.part"))
    assert external.read_bytes() == PDF


def test_new_profile_and_scratch_for_every_call(fake_engine):
    worker = LibreOfficeToPdfConverter()
    for _ in range(2):
        result = worker.normalize(NormalizationRequest(source=SOURCE, output_path=None))
        assert result.artifact.read_bytes() == PDF
    first = fake_engine[0][1]["cwd"]
    second = fake_engine[1][1]["cwd"]
    assert first != second
    assert not first.exists() and not second.exists()


@pytest.mark.parametrize("same_source", [False, True])
def test_existing_destination_and_source_not_overwritten(same_source, fake_engine, tmp_path):
    source = tmp_path / "original.docx"
    source.write_bytes(SOURCE)
    destination = source if same_source else tmp_path / "result.pdf"
    if not same_source:
        destination.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        LibreOfficeToPdfConverter().normalize(NormalizationRequest(source=source, output_path=destination))
    assert source.read_bytes() == SOURCE
    assert destination.read_bytes() == (SOURCE if same_source else b"existing")
    assert fake_engine == []
    assert not list(tmp_path.glob(".raggae-*.part"))


@pytest.mark.parametrize("write_result", [None, 0, -1, "error"])
def test_destination_errors_propagate_after_profile_cleanup(write_result, fake_engine):
    class BrokenOutput:
        def write(self, data):
            if write_result == "error":
                raise OSError("destination error")
            return write_result

    with BytesIO(SOURCE) as source, pytest.raises(OSError):
        LibreOfficeToPdfConverter()._transform(source, BrokenOutput())
    assert not fake_engine[0][1]["cwd"].exists()


def test_partial_writes_complete_pdf(fake_engine):
    class PartialOutput(BytesIO):
        def write(self, data):
            return super().write(data[:3])

    with BytesIO(SOURCE) as source, PartialOutput() as output:
        report = LibreOfficeToPdfConverter()._transform(source, output)
        assert output.getvalue() == PDF
        assert report.output_size_bytes == len(PDF)


def test_source_read_error_cleans_private_directory(fake_engine, monkeypatch):
    from tempfile import TemporaryDirectory

    scratch = []

    def directory(**kwargs):
        result = TemporaryDirectory(**kwargs)
        scratch.append(Path(result.name))
        return result

    class BrokenSource(BytesIO):
        def read(self, *args):
            raise OSError("source error")

    monkeypatch.setattr(libreoffice, "TemporaryDirectory", directory)
    with BrokenSource(SOURCE) as source, BytesIO() as output, pytest.raises(OSError, match="source error"):
        LibreOfficeToPdfConverter()._transform(source, output)
    assert fake_engine == []
    assert all(not path.exists() for path in scratch)


def test_scratch_read_error_is_storage_error(fake_engine, monkeypatch, tmp_path):
    original_read = Path.read_bytes

    def read(path):
        if path.name == "source.pdf":
            raise OSError("scratch read error")
        return original_read(path)

    monkeypatch.setattr(Path, "read_bytes", read)
    with pytest.raises(OSError, match="scratch read error"):
        LibreOfficeToPdfConverter().normalize(NormalizationRequest(source=SOURCE, output_path=tmp_path / "result.pdf"))
    assert not fake_engine[0][1]["cwd"].exists()
    assert not list(tmp_path.iterdir())


def test_publish_error_cleans_profile_and_staging(fake_engine, monkeypatch, tmp_path):
    from raggae.documents.normalization import io

    def broken_link(*args):
        raise OSError("publish error")

    monkeypatch.setattr(io.os, "link", broken_link)
    with pytest.raises(OSError, match="publish error"):
        LibreOfficeToPdfConverter().normalize(NormalizationRequest(source=SOURCE, output_path=tmp_path / "result.pdf"))
    assert not fake_engine[0][1]["cwd"].exists()
    assert not list(tmp_path.iterdir())


def test_worker_contract_and_unused_settings(fake_engine):
    worker = LibreOfficeToPdfConverter()
    assert isinstance(worker, NormalizationWorker)
    assert worker.name == "libreoffice_to_pdf_converter"
    assert not hasattr(worker, "implementation")
    assert not hasattr(worker, "version")
    with pytest.raises(TypeError, match="NormalizationRequest"):
        worker.normalize({})
    result = worker.normalize(NormalizationRequest(source=SOURCE, output_path=None, password="unused-secret", declared_encoding="ascii"))
    assert result.outcome == "changed"
    assert "unused-secret" not in repr(result)


@pytest.mark.parametrize("value,expected", [("DOCX", "docx"), (" .XLSX ", "xlsx"), (".PPTX", "pptx")])
def test_source_format_normalization(value, expected):
    assert LibreOfficeToPdfConverter(value).source_format == expected


@pytest.mark.parametrize("value", [None, False, 1, b"docx"])
def test_source_format_type_errors(value):
    with pytest.raises(TypeError):
        LibreOfficeToPdfConverter(value)


@pytest.mark.parametrize("value", ["", " ", "pdf", "odt", "doc", "markdown", "--convert-to=evil"])
def test_unsupported_routes_rejected(value):
    with pytest.raises(ValueError):
        LibreOfficeToPdfConverter(value)


@pytest.mark.parametrize("value", [False, True, None, 1.0, "120"])
def test_timeout_type_errors(value):
    with pytest.raises(TypeError):
        LibreOfficeToPdfConverter(timeout_seconds=value)


@pytest.mark.parametrize("value", [0, -1])
def test_timeout_must_be_positive(value):
    with pytest.raises(ValueError):
        LibreOfficeToPdfConverter(timeout_seconds=value)


@pytest.mark.parametrize("value", [False, 1, b"soffice"])
def test_executable_type_errors(value):
    with pytest.raises(TypeError):
        LibreOfficeToPdfConverter(executable=value)


@pytest.mark.parametrize("value", ["", " "])
def test_executable_cannot_be_blank(value):
    with pytest.raises(ValueError):
        LibreOfficeToPdfConverter(executable=value)


def test_constructor_does_not_discover_or_run_tools(monkeypatch):
    def forbidden(*args, **kwargs):
        pytest.fail("native discovery and execution must wait for normalize")

    monkeypatch.setattr(libreoffice.shutil, "which", forbidden)
    monkeypatch.setattr(libreoffice.subprocess, "run", forbidden)
    assert LibreOfficeToPdfConverter().source_format == "docx"


def test_report_immutable_and_serializable():
    report = OfficeToPdfReport(outcome="changed", source_format="docx", output_size_bytes=len(PDF))
    assert OfficeToPdfReport.model_validate_json(report.model_dump_json()) == report
    with pytest.raises(ValidationError):
        report.output_size_bytes = 1
    with pytest.raises(ValidationError):
        OfficeToPdfIssue(code=OfficeToPdfFinding.TIMEOUT, message=" ")


@pytest.mark.parametrize("kwargs", [
    {"outcome": "failed"}, {"outcome": "changed"}, {"outcome": "no_change"},
    {"outcome": "changed", "output_size_bytes": 0},
    {"outcome": "changed", "output_size_bytes": True},
    {"outcome": "changed", "output_size_bytes": "100"},
    {"outcome": "changed", "output_size_bytes": 1, "issues": [{"code": OfficeToPdfFinding.TIMEOUT, "message": "Timeout"}]},
    {"outcome": "failed", "output_size_bytes": 1, "issues": [{"code": OfficeToPdfFinding.TIMEOUT, "message": "Timeout"}]},
])
def test_inconsistent_reports_rejected(kwargs):
    with pytest.raises(ValidationError):
        OfficeToPdfReport(source_format="docx", **kwargs)


@pytest.mark.parametrize("source_format", ["docx", "xlsx", "pptx"])
def test_real_office_conversion(source_format):
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        pytest.skip("optional LibreOffice executable is not installed")
    # Real fixture builders are existing optional text-office dependencies,
    # not new runtime dependencies of this worker.
    with BytesIO() as encoded:
        if source_format == "docx":
            docx = pytest.importorskip("docx")
            document = docx.Document()
            document.add_paragraph("Normalization example")
            document.save(encoded)
        elif source_format == "xlsx":
            openpyxl = pytest.importorskip("openpyxl")
            workbook = openpyxl.Workbook()
            try:
                workbook.active["A1"] = "Normalization example"
                workbook.save(encoded)
            finally:
                workbook.close()
        else:
            pptx = pytest.importorskip("pptx")
            from pptx.util import Inches

            presentation = pptx.Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
            box.text_frame.text = "Normalization example"
            presentation.save(encoded)
        original = encoded.getvalue()
    request = NormalizationRequest(source=original, output_path=None)
    result = LibreOfficeToPdfConverter(source_format, executable=executable).normalize(request)
    assert result.outcome == "changed", result.report
    assert request.source.read_bytes() == original
    pdf = result.artifact.read_bytes()
    assert result.report.output_size_bytes == len(pdf)
    reader = PdfReader(BytesIO(pdf), strict=True)
    assert len(reader.pages) >= 1
    assert "Normalization example" in "\n".join(page.extract_text() for page in reader.pages)
